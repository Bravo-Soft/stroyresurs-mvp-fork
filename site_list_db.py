"""Лёгкий адаптер PostgreSQL-ситлиста для Диспетчерской.

Драйвер импортируется только в момент обращения: при отсутствии DSN или
``psycopg2`` панель остаётся доступной для наблюдения за прогонами и честно
возвращает 503 только на маршрутах ситлиста.
"""

from urllib.parse import unquote, urlsplit, urlunsplit


class SiteListUnavailable(RuntimeError):
    pass


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fetch_xlsx(dsn, company_id=None):
    path = unquote(dsn[len('xlsx://'):])
    if not path:
        raise SiteListUnavailable('Не задан путь к Site_list.xlsx')
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SiteListUnavailable('Не установлен openpyxl: XLSX-ситлист недоступен') from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [_cell_text(value) for value in next(rows, ())]
        positions = {name: index for index, name in enumerate(headers) if name}
        website_column = positions.get('Website')
        id_column = positions.get('ID производителя', positions.get('id производителя'))
        name_column = positions.get('Наименование', positions.get('Название'))
        if website_column is None:
            raise SiteListUnavailable('В Site_list.xlsx нет обязательной колонки Website')
        result = []
        for index, row in enumerate(rows):
            website = _cell_text(row[website_column] if website_column < len(row) else None)
            if not website:
                continue
            name = _cell_text(row[name_column] if name_column is not None and name_column < len(row) else None)
            identifier = _cell_text(row[id_column] if id_column is not None and id_column < len(row) else None)
            item = {'company_id': identifier or f'company_{len(result)}',
                    'name': name, 'website': website}
            if company_id is None or item['company_id'] == str(company_id):
                result.append(item)
        workbook.close()
        return result
    except SiteListUnavailable:
        raise
    except Exception as exc:
        raise SiteListUnavailable(f'Не удалось прочитать XLSX-ситлист: {exc}') from exc


def _connect(dsn):
    if not dsn:
        raise SiteListUnavailable('SITE_LIST_DSN не задан: ситлист PostgreSQL недоступен')
    try:
        import psycopg2
    except ImportError as exc:
        raise SiteListUnavailable('Не установлен psycopg2-binary: ситлист PostgreSQL недоступен') from exc
    try:
        return psycopg2.connect(dsn, connect_timeout=5)
    except Exception as exc:
        raise SiteListUnavailable(f'Не удалось подключиться к ситлисту PostgreSQL: {exc}') from exc


def describe_dsn(dsn):
    """Адрес без пароля, пригодный для API-ответа и журнала."""
    if not dsn:
        return 'не настроен'
    if str(dsn).startswith('xlsx://'):
        return 'XLSX: ' + unquote(str(dsn)[len('xlsx://'):])
    try:
        parsed = urlsplit(dsn)
        if parsed.scheme and parsed.hostname:
            host = parsed.hostname + (f':{parsed.port}' if parsed.port else '')
            return urlunsplit((parsed.scheme, host, parsed.path, '', ''))
    except (TypeError, ValueError):
        pass
    return '<настроен>'


def _columns(connection):
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT column_name FROM information_schema.columns "
            "WHERE table_schema = 'public' AND table_name = 'companies'"
        )
        return {row[0] for row in cursor.fetchall()}


def _choose(columns, candidates, label):
    for name in candidates:
        if name in columns:
            return name
    raise SiteListUnavailable(
        f'В public.companies не найдена колонка {label}; доступны: {", ".join(sorted(columns))}'
    )


def _query_columns(connection):
    columns = _columns(connection)
    return (
        _choose(columns, ('company_id', 'manufacturer_id', 'id'), 'идентификатора компании'),
        _choose(columns, ('name', 'company_name', 'manufacturer_name', 'original_name'), 'названия'),
        _choose(columns, ('website', 'url', 'site_url'), 'адреса сайта'),
    )


def _fetch(dsn, company_id=None):
    if str(dsn or '').startswith('xlsx://'):
        return _fetch_xlsx(dsn, company_id)
    connection = _connect(dsn)
    try:
        id_column, name_column, website_column = _query_columns(connection)
        sql = (
            f'SELECT "{id_column}"::text, "{name_column}", "{website_column}" '
            'FROM public.companies'
        )
        params = []
        if company_id is not None:
            sql += f' WHERE "{id_column}"::text = %s'
            params.append(str(company_id))
        sql += f' ORDER BY "{id_column}"'
        with connection.cursor() as cursor:
            cursor.execute(sql, params)
            rows = cursor.fetchall()
        return [
            {'company_id': str(row[0]), 'name': row[1] or '', 'website': row[2] or ''}
            for row in rows
        ]
    except SiteListUnavailable:
        raise
    except Exception as exc:
        raise SiteListUnavailable(f'Не удалось прочитать public.companies: {exc}') from exc
    finally:
        connection.close()


def fetch_all(dsn):
    return _fetch(dsn)


def fetch_one(dsn, company_id):
    rows = _fetch(dsn, company_id)
    return rows[0] if rows else None
