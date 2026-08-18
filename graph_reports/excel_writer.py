# excel_writer.py 1.0.0
import os
import logging
from typing import List, Any
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("excel_writer")

class ExcelWriter:
    """Класс для записи данных в Excel"""
    
    def __init__(self, config):
        self.config = config
    
    def write_general_report(self, data: List[List[Any]], output_path: str = None) -> str:
        """Запись общего отчета в Excel"""
        if not output_path:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = os.path.join(
                self.config.output_dir, 
                f"general_report_{timestamp}.xlsx"
            )
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Отчёт"
        
        # Записываем данные
        for row_idx, row in enumerate(data, 1):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Форматирование
        self._apply_general_report_formatting(ws, len(data[0]) if data else 0)
        
        # Убираем автоподбор для общего отчета, так как ширина задана вручную
        # self._auto_adjust_columns(ws)
        
        # Сохраняем файл
        wb.save(output_path)
        log.info(f"Общий отчет сохранен: {output_path}")
        
        return output_path
    
    def write_detailed_report(self, data: List[List[Any]], output_path: str = None) -> str:
        """Запись детального отчет в Excel"""
        if not output_path:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = os.path.join(
                self.config.output_dir, 
                f"detailed_report_{timestamp}.xlsx"
            )
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "отчет"
        
        # Записываем данные
        for row_idx, row in enumerate(data, 1):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Форматирование
        self._apply_detailed_report_formatting(ws, len(data[0]) if data else 0)
        
        # Автоподбор ширины колонок (оставляем для детального отчета)
        self._auto_adjust_columns(ws)
        
        # Сохраняем файл
        wb.save(output_path)
        log.info(f"Детальный отчет сохранен: {output_path}")
        
        return output_path
    
    def _apply_general_report_formatting(self, ws, num_columns: int):
        """Применение форматирования для общего отчета"""
        # Заголовки
        header_fill = PatternFill(start_color="224681", end_color="224681", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        for col in range(1, num_columns + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Установка ширины колонок (13 столбцов) В РУЧНУЮ
        column_widths = {
            "A": 12,  # Дата проверки сайта (upload_date/update_date)
            "B": 20,  # Наличие изменений (4 статуса)
            "C": 12,  # ID компании
            "D": 40,  # Наименование компании
            "E": 50,  # Адрес
            "F": 25,  # Телефоны
            "G": 30,  # Сайт
            "H": 25,  # e-mail
            "I": 20,  # Регион
            "J": 80,  # Информация о компании
            "K": 17,  # Количество поставщиков
            "L": 17,  # Количество товаров
            "M": 17,  # Количество файлов производителя
        }
        
        for col_letter, width in column_widths.items():
            # Проверяем, существует ли такая колонка в листе
            col_idx = openpyxl.utils.column_index_from_string(col_letter)
            if col_idx <= num_columns:
                ws.column_dimensions[col_letter].width = width
        
        # Включаем перенос текста для колонки с информацией
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=10).alignment = Alignment(wrap_text=True, vertical="top")
        
        # Добавляем цветовое кодирование для статусов
        self._apply_status_formatting(ws)
        
        # Добавляем специальное форматирование для компаний без представительства
        self._apply_no_representation_formatting(ws)
        
        # ДОБАВЛЕНО: Применение рамок ко всем ячейкам
        self._apply_borders_to_all_cells(ws, ws.max_row, num_columns)
    
    def _apply_detailed_report_formatting(self, ws, num_columns: int):
        """Применение форматирования для детального отчета"""
        # Заголовки
        header_fill = PatternFill(start_color="224681", end_color="224681", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        for col in range(1, num_columns + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Установка ширины колонок - уменьшаем ширину для пустых ID столбцов
        column_widths = {
            "A": 10,  # п/п изм.
            "B": 15,  # ID производителя
            "C": 40,  # Название компании
            "D": 10,  # ID поставщика (теперь пустое, уменьшаем ширину)
            "E": 40,  # Наименование поставщика
            "F": 10,  # ID товара (теперь пустое, уменьшаем ширину)
            "G": 60,  # Наименование товара
            "H": 10,  # ID доп.материала (теперь пустое, уменьшаем ширину)
            "I": 30,  # Наименование доп. материала
            "J": 40,  # Было
            "K": 40,  # Стало
            "L": 20,  # Тип изменения
            "M": 40,  # Ссылка на источник
            "N": 80   # Путь в архиве (теперь без папки компании)
        }
        
        for col_letter, width in column_widths.items():
            # Проверяем, существует ли такая колонка в листе
            col_idx = openpyxl.utils.column_index_from_string(col_letter)
            if col_idx <= num_columns:
                ws.column_dimensions[col_letter].width = width
        
        # Добавляем форматирование для разных типов изменений
        self._apply_change_type_formatting(ws)
        
        # ДОБАВЛЕНО: Применение рамок ко всем ячейкам
        self._apply_borders_to_all_cells(ws, ws.max_row, num_columns)

    def _apply_borders_to_all_cells(self, ws, max_row: int, max_col: int):
        """Применение рамок ко всем ячейкам в отчете"""
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Применяем рамки ко всем ячейкам от A1 до последней заполненной
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.border = thin_border

    def _apply_status_formatting(self, ws):
        """Применение форматирования для статусов в общем отчете"""
        # Колонка B - Наличие изменений
        status_col = 2
        
        # Цвета для разных статусов
        status_colors = {
            "Новая компания": "C6EFCE",  # Светло-зеленый
            "Есть изменения": "FFEB9C",  # Светло-желтый
            "Нет изменений": "FFFFFF",   # Белый
            "Сайт не найден": "CFBDBD"   # Серо-красный
        }
        
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=status_col)
            status = str(cell.value).strip()
            
            if status in status_colors:
                fill = PatternFill(start_color=status_colors[status], 
                                 end_color=status_colors[status], 
                                 fill_type="solid")
                cell.fill = fill
    
    def _apply_change_type_formatting(self, ws):
        """Применение форматирования в зависимости от типа изменения"""
        # Определяем индекс колонки с типом изменения (колонка L)
        change_type_col = 12  # Колонка L
        
        # Цвета для разных типов изменений
        color_map = {
            "Новые данные": "C6EFCE",  # Светло-зеленый
            "Данные не найдены": "F07575",  # Светло-красный
            "Редактирование": "FFEB9C",  # Светло-желтый
            "Ошибка, сайт не работает": "CFBDBD",  # Серо-красный
            "Изменений не найдено": "FFFFFF",  # Белый
            "Новая компания": "C6EFCE"  # Светло-зеленый
        }
        
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=change_type_col)
            change_type = str(cell.value).strip()
            
            if change_type in color_map:
                fill = PatternFill(start_color=color_map[change_type], 
                                 end_color=color_map[change_type], 
                                 fill_type="solid")
                cell.fill = fill
            
            # Перенос текста для колонок "Было" и "Стало"
            ws.cell(row=row, column=10).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=11).alignment = Alignment(wrap_text=True, vertical="top")
    
    def _apply_no_representation_formatting(self, ws):
        """Применение форматирования для компаний без представительства"""
        # Колонка E - Адрес компании производителя (столбец 5)
        address_col = 5
        
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=address_col)
            address_value = str(cell.value).strip()
            
            if address_value == "Представительство на территории РФ отсутствует":
                # Применяем серый фон для всей строки
                gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
                for col in range(1, 14):  # 13 столбцов
                    ws.cell(row=row, column=col).fill = gray_fill
                
                # Жирный курсив для адреса
                cell.font = Font(italic=True, bold=True, color="FF0000")
    
    def _auto_adjust_columns(self, ws):
        """Автоподбор ширины колонок - используется только для детального отчета"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            # Проверяем, не была ли уже установлена ширина вручную
            current_width = ws.column_dimensions[column_letter].width
            if current_width is not None:
                # Если ширина уже установлена (не None), пропускаем автоподбор для этой колонки
                continue
                
            for cell in column:
                try:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 100)  # Максимальная ширина 100
            ws.column_dimensions[column_letter].width = adjusted_width