# xlsx_generator.py 1.0.0
import pandas as pd
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import logging
from datetime import datetime
import os
from config import Config

log = logging.getLogger("xlsx_generator")

class XLSXGenerator:
    """Генератор отчетов в формате XLSX"""
    def __init__(self, config):
        self.config = config
        os.makedirs(self.config.reports_dir, exist_ok=True)
    
    def generate_monitoring_report(self, results: List[Dict[str, Any]], statistics: Any, reports_dir: str, timestamp: str) -> str:
        """Генерация отчета мониторинга в формате XLSX"""
        
        # Создаем Excel writer
        output_path = os.path.join(reports_dir, f"monitoring_report_{timestamp}.xlsx")
        os.makedirs(reports_dir, exist_ok=True)
        
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # Лист с общей статистикой
            self._create_summary_sheet(writer, statistics, results)
            
            # Лист с детализацией по компаниям
            self._create_companies_sheet(writer, results, statistics)
            
            # Лист со статистикой по токенам
            self._create_tokens_sheet(writer, results, statistics)

            # Лист со статистикой по времени обработки
            self._create_processing_time_sheet(writer, results, statistics)
        
        log.info(f"XLSX отчет сохранен: {output_path}")
        return output_path
    
    def _create_summary_sheet(self, writer, statistics, results: List[Dict[str, Any]]):
        """Создание листа с общей статистикой"""
        total_products_processed = sum(r.get('company_statistics', {}).get('products_processed', 0) for r in results)
        summary_data = {
            'Показатель': [
                'Количество сайтов для проверки',
                'Успешно обработано сайтов',
                'Ошибки обработки сайтов',
                'Страниц обработано',
                'Продуктовых страниц обработано',
                'Страниц компаний обработано',
                'Страниц дистрибьюторов обработано',
                'Товаров обработано (полная обработка)',
                'Товаров отправлено в граф БД',
                'Файлов найдено',
                'Файлов скачано',
                'Файлов с ошибкой скачивания',
                'Всего входных токенов',
                'Всего выходных токенов',
                'Общее количество токенов',
                'Общее время выполнения'
            ],
            'Значение': [
                len(results),
                len([r for r in results if r.get('status') == 'success']),
                len([r for r in results if r.get('status') != 'success']),
                statistics.pages_processed,
                statistics.product_pages_processed,
                statistics.company_pages_processed,
                statistics.distributor_pages_processed,
                total_products_processed,
                statistics.graph_db_products_sent,
                statistics.total_files_found,
                statistics.total_files_downloaded,
                statistics.total_files_failed,
                statistics.total_input_tokens,
                statistics.total_output_tokens,
                statistics.total_input_tokens + statistics.total_output_tokens,
                statistics.get_execution_time()
            ]
        }
        
        df_summary = pd.DataFrame(summary_data)
        df_summary.to_excel(writer, sheet_name='Общая статистика', index=False)
        
        # Форматирование
        workbook = writer.book
        worksheet = writer.sheets['Общая статистика']
        
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    def _create_companies_sheet(self, writer, results: List[Dict[str, Any]], statistics):
        """Создание листа с детализацией по компаниям"""
        companies_data = []
        
        for result in results:
            company_data = result.get('company', {})
            company_stats = result.get('company_statistics', {})
            
            # Собираем ошибки
            errors = result.get('errors', [])
            error_text = "; ".join(errors) if errors else ""
            
            companies_data.append({
                'ID производителя': company_data.get('company_id', ''),
                'Наименование компании': company_data.get('original_name', ''),
                'URL сайта': company_data.get('website', ''),
                'Ошибка при обработке сайта': error_text,
                'Товаров обработано (полная обработка)': company_stats.get('products_processed', 0),
                'Не удалось обработать': company_stats.get('products_failed', 0),
                'Продуктовых страниц': company_stats.get('product_pages_processed', 0),
                'Страниц компании': company_stats.get('company_pages_processed', 0),
                'Страниц дистрибьюторов': company_stats.get('distributor_pages_processed', 0),
                'Найдено файлов для скачивания': company_stats.get('files_found', 0),
                'Успешно скачаны': company_stats.get('files_downloaded', 0),
                'Ошибка скачивания': company_stats.get('files_failed', 0),
                'Добавлено товаров в базу данных': company_stats.get('products_added_to_db', 0),
                'Сформировано карточек товаров': company_stats.get('cards_generated', 0),
                'Конвертировано файлов в PDF': company_stats.get('files_converted', 0),
                'Ошибка конвертации': company_stats.get('files_conversion_failed', 0)
            })
        
        df_companies = pd.DataFrame(companies_data)
        df_companies.to_excel(writer, sheet_name='Детализация по компаниям', index=False)
        
        # Форматирование
        worksheet = writer.sheets['Детализация по компаниям']
        
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            worksheet.column_dimensions[column_letter].width = adjusted_width

    def _create_tokens_sheet(self, writer, results: List[Dict[str, Any]], statistics):
        """Создание листа со статистикой по затраченным токенам"""
        tokens_data = []

        for result in results:
            company_data = result.get('company', {})
            cs = result.get('company_statistics', {})

            input_tokens = cs.get('input_tokens', 0)
            output_tokens = cs.get('output_tokens', 0)

            tokens_data.append({
                'ID производителя': company_data.get('company_id', ''),
                'Наименование компании': company_data.get('original_name', ''),
                'Входные токены': input_tokens,
                'Выходные токены': output_tokens,
                'Токены LLM (всего)': input_tokens + output_tokens
            })

        # Итоговая строка по всем компаниям
        tokens_data.append({
            'ID производителя': '',
            'Наименование компании': 'ИТОГО',
            'Входные токены': statistics.total_input_tokens,
            'Выходные токены': statistics.total_output_tokens,
            'Токены LLM (всего)': statistics.total_input_tokens + statistics.total_output_tokens
        })

        df_tokens = pd.DataFrame(tokens_data)
        df_tokens.to_excel(writer, sheet_name='Статистика по токенам', index=False)

        # Форматирование
        worksheet = writer.sheets['Статистика по токенам']
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            worksheet.column_dimensions[column_letter].width = adjusted_width


    def _create_processing_time_sheet(self, writer, results: List[Dict[str, Any]], statistics):
        """Лист «Статистика по времени обработки» — строка на компанию + ИТОГО.
        Чистое время обработки единицы каждым модулем (Вариант A) и время ретраев.
        Данные берутся из statistics.processing_tracker (если нет — нули)."""
        try:
            from processing_time_tracker import MODULE_ORDER, MODULE_LABELS, RETRY_MODULES
        except Exception:
            MODULE_ORDER, MODULE_LABELS, RETRY_MODULES = [], {}, set()

        tracker = getattr(statistics, 'processing_tracker', None)
        _empty = {'count': 0, 'total': 0.0, 'median': 0.0, 'retry': 0.0}

        def _fill(row, module, summ):
            label, unit = MODULE_LABELS.get(module, (module, 'ед.'))
            row[f"{label}: {unit}"] = summ['count']
            row[f"{label}: чистое время, с"] = round(summ['total'], 2)
            row[f"{label}: медиана/ед., с"] = round(summ['median'], 3)
            if module in RETRY_MODULES:
                row[f"{label}: ретраи, с"] = round(summ['retry'], 2)

        rows = []
        for result in results:
            company_data = result.get('company', {})
            cid = company_data.get('company_id', '')
            row = {
                'ID производителя': cid,
                'Наименование компании': company_data.get('original_name', ''),
            }
            for module in MODULE_ORDER:
                summ = tracker.company_summary(cid, module) if tracker else _empty
                _fill(row, module, summ)
            rows.append(row)

        total_row = {'ID производителя': '', 'Наименование компании': 'ИТОГО'}
        for module in MODULE_ORDER:
            summ = tracker.global_summary(module) if tracker else _empty
            _fill(total_row, module, summ)
        rows.append(total_row)

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='Статистика по времени обработки', index=False)

        worksheet = writer.sheets['Статистика по времени обработки']
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 26)

    def generate_detailed_changes_report(self, changes_data: Dict[str, Any], 
                                       multibrand_data: Dict[str, Any],
                                       timestamp: str) -> str:
        """Генерация детального отчета об изменениях в формате XLSX"""
        try:
            output_path = os.path.join(self.config.reports_dir, f"Детальный_отчет_изменений_{timestamp}.xlsx")
            
            workbook = Workbook()
            workbook.remove(workbook.active)  # Удаляем дефолтный лист
            
            # Создаем лист с детальными изменениями
            detailed_sheet = workbook.create_sheet("Детальный отчет изменений")
            
            # Заголовки столбцов
            headers = [
                "п/п изм.",
                "ID производителя", 
                "Название компании производителя",
                "ID товара",
                "Наименование товара", 
                "ID доп.материала",
                "Наименование доп. материала",
                "Было (в нашей БД)",
                "Стало (на сайте)",
                "Тип изменения",
                "Ссылка на источник",
                "Путь в архиве к файлу"
            ]
            
            # Добавляем заголовки
            detailed_sheet.append(headers)
            
            # Форматируем заголовки
            self._format_header_row(detailed_sheet[1])
            
            # Заполняем данные
            row_num = 2
            change_counter = 1
            
            # Обрабатываем изменения компаний
            for company_id, change_info in changes_data.get('companies', {}).items():
                change_row = self._create_company_change_row(change_info, company_id, multibrand_data, change_counter)
                if change_row:
                    detailed_sheet.append(change_row)
                    self._format_data_row(detailed_sheet[row_num], change_info.get('change_type', ''))
                    row_num += 1
                    change_counter += 1
            
            # Обрабатываем изменения продуктов
            for product_id, change_info in changes_data.get('products', {}).items():
                change_row = self._create_product_change_row(change_info, multibrand_data, change_counter)
                if change_row:
                    detailed_sheet.append(change_row)
                    self._format_data_row(detailed_sheet[row_num], change_info.get('change_type', ''))
                    row_num += 1
                    change_counter += 1
            
            # Обрабатываем изменения файлов
            for file_id, change_info in changes_data.get('files', {}).items():
                change_row = self._create_file_change_row(change_info, multibrand_data, change_counter)
                if change_row:
                    detailed_sheet.append(change_row)
                    self._format_data_row(detailed_sheet[row_num], change_info.get('change_type', ''))
                    row_num += 1
                    change_counter += 1
            
            # Автоподбор ширины столбцов
            self._auto_adjust_columns(detailed_sheet)
            
            workbook.save(output_path)
            log.info(f"Детальный отчет об изменениях сохранен: {output_path}")
            return output_path
            
        except Exception as e:
            log.error(f"Ошибка генерации детального отчета об изменениях: {e}")
            return ""
    
    def generate_summary_changes_report(self, current_data: Dict[str, Any],
                                      changes_data: Dict[str, Any], 
                                      multibrand_data: Dict[str, Any],
                                      timestamp: str) -> str:
        """Генерация сводного отчета об изменениях в формате XLSX"""
        try:
            output_path = os.path.join(self.config.reports_dir, f"Общий_отчет_изменений_{timestamp}.xlsx")
            
            workbook = Workbook()
            workbook.remove(workbook.active)  # Удаляем дефолтный лист
            
            # Создаем лист с общим отчетом
            summary_sheet = workbook.create_sheet("Общий отчет изменений")
            
            # Заголовки столбцов
            headers = [
                "Дата проверки сайта",
                "Наличие изменений на сайте компании производителя", 
                "ID компании производителя",
                "Наименование компании производителя",
                "Адрес компании производителя",
                "Телефоны компании производителя", 
                "Сайт компании производителя",
                "e-mail компании производителя",
                "Регион производства",
                "Информация о компании производителе",
                "ID материалов компании производителя", 
                "Путь в архиве к файлу компании",
                "ID компании поставщика",
                "Наименование компании поставщика",
                "Адрес компании поставщика",
                "Телефоны компании поставщика", 
                "Сайт компании поставщика",
                "E-mail",
                "Регион поставщика", 
                "Информация о компании поставщике",
                "ID товара",
                "Название товара",
                "Путь к файлу с описанием товара", 
                "ID дополнительных материалов к товарам",
                "Путь к дополнительным файлам товаров"
            ]
            
            # Добавляем заголовки
            summary_sheet.append(headers)
            
            # Форматируем заголовки
            self._format_header_row(summary_sheet[1])
            
            # Заполняем данные
            row_num = 2
            
            for company_id, company_data in current_data.items():
                company_rows = self._create_company_summary_rows(company_data, changes_data, multibrand_data)
                for row in company_rows:
                    summary_sheet.append(row)
                    
                    # Определяем цвет строки в зависимости от наличия изменений
                    has_changes = self._company_has_changes(company_id, changes_data)
                    change_type = self._get_company_change_type(company_id, changes_data)
                    self._format_summary_data_row(summary_sheet[row_num], has_changes, change_type)
                    
                    row_num += 1
            
            # Автоподбор ширины столбцов
            self._auto_adjust_columns(summary_sheet)
            
            workbook.save(output_path)
            log.info(f"Сводный отчет об изменениях сохранен: {output_path}")
            return output_path
            
        except Exception as e:
            log.error(f"Ошибка генерации сводного отчета об изменениях: {e}")
            return ""
    
    def _create_company_change_row(self, change_info: Dict[str, Any], company_id: str, 
                                 multibrand_data: Dict[str, Any], change_counter: int) -> List:
        """Создание строки изменения компании для детального отчета"""
        change_type = change_info.get('change_type', '')
        current_data = change_info.get('current', {})
        previous_data = change_info.get('previous', {})
        details = change_info.get('details', [])
        
        # Получаем информацию о компании из multibrand_data
        company_info = multibrand_data.get(company_id, {}).get('company_info', {})
        company_name = company_info.get('name', '')
        if not company_name:
            # Если нет в multibrand_data, берем из current_data или previous_data
            company_data = current_data.get('company', {}) if current_data else previous_data.get('company', {})
            company_name = company_data.get('original_name', '')
        
        # Определяем значения для "Было" и "Стало" в зависимости от типа изменения
        if change_type == "Новые данные":
            was_value = "Данные в базе отсутствуют"
            became_value = "; ".join(details) if details else "Новые данные компании"
        elif change_type == "Данные не найдены":
            was_value = "; ".join(details) if details else "Данные компании"
            became_value = "Данные на сайте отсутствуют"
        elif change_type == "Редактирование":
            was_value = " | ".join([f"{detail}" for detail in details]) if details else "Предыдущие данные"
            became_value = " | ".join([f"{detail}" for detail in details]) if details else "Новые данные"
        elif change_type == "Ошибка, сайт не работает":
            was_value = ""
            became_value = ""
        elif change_type == "Изменений не найдено":
            was_value = ""
            became_value = ""
        else:
            was_value = ""
            became_value = ""
        
        # Определяем ссылку на источник
        website = current_data.get('company', {}).get('website', '') if current_data else previous_data.get('company', {}).get('website', '')
        source_link = website if website else ""
        
        if change_type == "Ошибка, сайт не работает":
            source_link = website  # Сохраняем URL проблемного сайта
        
        # Формируем путь к файлам компании
        company_file_path = self._generate_company_documents_path(company_id, company_name)
        
        return [
            change_counter,      # п/п изм.
            company_id,          # ID производителя
            company_name,        # Название компании производителя
            "",                  # ID товара
            "",                  # Наименование товара
            "",                  # ID доп.материала
            "",                  # Наименование доп. материала
            was_value,           # Было (в нашей БД)
            became_value,        # Стало (на сайте)
            change_type,         # Тип изменения
            source_link,         # Ссылка на источник
            company_file_path    # Путь в архиве к файлу
        ]
    
    def _create_product_change_row(self, change_info: Dict[str, Any], multibrand_data: Dict[str, Any], change_counter: int) -> List:
        """Создание строки изменения продукта для детального отчета"""
        change_type = change_info.get('change_type', '')
        current_data = change_info.get('current', {})
        previous_data = change_info.get('previous', {})
        details = change_info.get('details', [])
        
        product_data = current_data.get('data', {}) if current_data else previous_data.get('data', {})
        product_id = current_data.get('product_id', '') if current_data else previous_data.get('product_id', '')
        company_id = current_data.get('company_id', '') if current_data else previous_data.get('company_id', '')
        
        # Получаем полный URL страницы товара
        source_url = current_data.get('source_url', '') if current_data else previous_data.get('source_url', '')
        if not source_url:
            # Пытаемся получить из metadata
            source_url = current_data.get('metadata', {}).get('url', '') if current_data else previous_data.get('metadata', {}).get('url', '')
        
        product_name = product_data.get('name', '')
        
        # Получаем информацию о компании из multibrand_data
        company_info = multibrand_data.get(company_id, {}).get('company_info', {})
        company_name = company_info.get('name', '')
        if not company_name:
            # Если нет в multibrand_data, берем из данных товара
            company_name = current_data.get('manufacturer', '') if current_data else previous_data.get('manufacturer', '')
        
        # Определяем значения для "Было" и "Стало"
        if change_type == "Новые данные":
            was_value = "Данные в базе отсутствуют"
            became_value = product_name
        elif change_type == "Данные не найдены":
            was_value = product_name
            became_value = "Данные на сайте отсутствуют"
        elif change_type == "Редактирование":
            was_value = " | ".join(details) if details else "Предыдущие данные товара"
            became_value = " | ".join(details) if details else "Новые данные товара"
        else:
            was_value = ""
            became_value = ""
        
        # Формируем пути к файлам товара
        product_main_path = self._generate_product_main_path(company_id, company_name, product_name)
        product_documents_path = self._generate_product_documents_path(company_id, company_name, product_name)
        
        return [
            change_counter,      # п/п изм.
            company_id,          # ID производителя
            company_name,        # Название компании производителя
            product_id,          # ID товара
            product_name,        # Наименование товара
            "",                  # ID доп.материала
            "",                  # Наименование доп. материала
            was_value,           # Было (в нашей БД)
            became_value,        # Стало (на сайте)
            change_type,         # Тип изменения
            source_url,          # Ссылка на источник (полный URL)
            product_main_path    # Путь в архиве к файлу
        ]
    
    def _create_file_change_row(self, change_info: Dict[str, Any], multibrand_data: Dict[str, Any], change_counter: int) -> List:
        """Создание строки изменения файла для детального отчета"""
        change_type = change_info.get('change_type', '')
        current_data = change_info.get('current', {})
        previous_data = change_info.get('previous', {})
        details = change_info.get('details', [])
        
        file_data = current_data if current_data else previous_data
        file_id = file_data.get('file_id', '')
        file_name = file_data.get('file_name', '')
        file_path = file_data.get('file_path', '')
        
        # Пытаемся получить информацию о товаре и компании из file_data
        product_id = file_data.get('product_id', '')
        company_id = file_data.get('company_id', '')
        
        # Получаем информацию о компании из multibrand_data
        company_name = ""
        if company_id:
            company_info = multibrand_data.get(company_id, {}).get('company_info', {})
            company_name = company_info.get('name', '')
        
        # Получаем информацию о товаре для формирования пути
        product_name = ""
        if product_id and company_id and company_name:
            # Здесь нужно получить название товара по product_id
            # В реальной системе это может быть поиск в данных товаров
            # Пока оставляем пустым, путь будет сформирован без названия товара
            pass
        
        # Определяем значения для "Было" и "Стало"
        if change_type == "Новые данные":
            was_value = "Данные в базе отсутствуют"
            became_value = file_name
        elif change_type == "Данные не найдены":
            was_value = file_name
            became_value = "Данные на сайте отсутствуют"
        elif change_type == "Редактирование":
            was_value = " | ".join(details) if details else "Предыдущая версия файла"
            became_value = " | ".join(details) if details else "Новая версия файла"
        else:
            was_value = ""
            became_value = ""
        
        # Формируем путь к файлу
        file_storage_path = ""
        if company_id and company_name:
            if product_id and product_name:
                # Файл товара
                file_storage_path = self._generate_product_documents_path(company_id, company_name, product_name)
            else:
                # Файл компании
                file_storage_path = self._generate_company_documents_path(company_id, company_name)
        
        return [
            change_counter,      # п/п изм.
            company_id,          # ID производителя
            company_name,        # Название компании производителя
            "",                  # ID товара
            "",                  # Наименование товара
            file_id,             # ID доп.материала
            file_name,           # Наименование доп. материала
            was_value,           # Было (в нашей БД)
            became_value,        # Стало (на сайте)
            change_type,         # Тип изменения
            file_data.get('original_url', ''),  # Ссылка на источник
            file_storage_path    # Путь в архиве к файлу
        ]
    
    def _create_company_summary_rows(self, company_data: Dict[str, Any], 
                                   changes_data: Dict[str, Any],
                                   multibrand_data: Dict[str, Any]) -> List[List]:
        """Создание строк сводного отчета для компании (ВСЕХ товаров)"""
        rows = []
        
        company_info = company_data.get('company', {})
        company_id = company_info.get('company_id', '')
        company_name = company_info.get('original_name', '')
        
        # Проверяем наличие изменений для компании
        has_changes = self._company_has_changes(company_id, changes_data)
        change_status = "Есть изменения" if has_changes else "Нет изменений"
        
        # Получаем информацию о мультибрендовости
        is_multibrand = multibrand_data.get(company_id, {}).get('is_multibrand', False)
        websites_data = multibrand_data.get(company_id, {}).get('websites_data', {})
        
        if is_multibrand:
            # Для мультибрендовых компаний создаем строку для каждого сайта
            for website, website_data in websites_data.items():
                # Получаем товары для этого сайта
                site_products = self._get_products_for_website(company_data, website)
                for product in site_products:
                    row = self._create_product_summary_row(company_info, change_status, website, product)
                    rows.append(row)
                
                # Если нет товаров для этого сайта, создаем строку без товара
                if not site_products:
                    row = self._create_product_summary_row(company_info, change_status, website, None)
                    rows.append(row)
        else:
            # Для обычных компаний - все товары
            website = company_info.get('website', '')
            products = company_data.get('products', [])
            
            for product in products:
                row = self._create_product_summary_row(company_info, change_status, website, product)
                rows.append(row)
            
            # Если нет товаров, создаем одну строку без товара
            if not products:
                row = self._create_product_summary_row(company_info, change_status, website, None)
                rows.append(row)
        
        return rows
    
    def _create_product_summary_row(self, company_info: Dict[str, Any],
                                  change_status: str, 
                                  website: str, 
                                  product: Optional[Dict[str, Any]]) -> List:
        """Создание одной строки сводного отчета для товара"""
        company_id = company_info.get('company_id', '')
        company_name = company_info.get('original_name', '')
        
        # Основная информация о компании
        row = [
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),  # Дата проверки сайта
            change_status,                                  # Наличие изменений
            company_id,                                     # ID компании производителя
            company_name,                                   # Наименование компании производителя
            company_info.get('address', ''),               # Адрес компании производителя
            company_info.get('phones', ''),                # Телефоны компании производителя
            website,                                       # Сайт компании производителя
            company_info.get('email', ''),                 # e-mail компании производителя
            company_info.get('region', ''),                # Регион производства
            company_info.get('info', ''),                  # Информация о компании производителе
            "",                                            # ID материалов компании производителя
            self._generate_company_documents_path(company_id, company_name),  # Путь в архиве к файлу компании
        ]
        
        # Добавляем поставщиков (в данном примере - заглушка)
        row.extend([
            "",  # ID компании поставщика
            "",  # Наименование компании поставщика  
            "",  # Адрес компании поставщика
            "",  # Телефоны компании поставщика
            "",  # Сайт компании поставщика
            "",  # E-mail
            "",  # Регион поставщика
            "",  # Информация о компании поставщике
        ])
        
        # Добавляем информацию о товаре, если он есть
        if product:
            product_data = product.get('data', {})
            product_name = product_data.get('name', 'Неизвестный товар')
            row.extend([
                product.get('product_id', ''),       # ID товара
                product_name,                       # Название товара
                self._generate_product_main_path(company_id, company_name, product_name),  # Путь к файлу с описанием товара
                "",                                 # ID дополнительных материалов к товарам
                self._generate_product_documents_path(company_id, company_name, product_name)  # Путь к дополнительным файлам товаров
            ])
        else:
            row.extend(["", "", "", "", ""])
        
        return row
    
    def _get_products_for_website(self, company_data: Dict[str, Any], website: str) -> List[Dict[str, Any]]:
        """Получение товаров для конкретного сайта (для мультибрендовых компаний)"""
        # В текущей реализации все товары компании считаются принадлежащими всем сайтам
        # В реальной системе здесь должна быть логика фильтрации товаров по сайту
        return company_data.get('products', [])
    
    def _generate_company_documents_path(self, company_id: str, company_name: str) -> str:
        """Генерация пути к документам компании согласно новой структуре"""
        safe_name = "".join(c for c in company_name if c.isalnum() or c in (' ', '_', '-')).rstrip()
        return f"{company_id}_{safe_name}\\Documents"
    
    def _generate_product_main_path(self, company_id: str, company_name: str, product_name: str) -> str:
        """Генерация пути к основным файлам товара согласно новой структуре"""
        safe_company_name = "".join(c for c in company_name if c.isalnum() or c in (' ', '_', '-')).rstrip()
        safe_product_name = "".join(c for c in product_name if c.isalnum() or c in (' ', '_', '-')).rstrip()
        return f"{company_id}_{safe_company_name}\\Products\\{safe_product_name}"
    
    def _generate_product_documents_path(self, company_id: str, company_name: str, product_name: str) -> str:
        """Генерация пути к документам товара согласно новой структуре"""
        product_path = self._generate_product_main_path(company_id, company_name, product_name)
        return f"{product_path}\\Documents"
    
    def _company_has_changes(self, company_id: str, changes_data: Dict[str, Any]) -> bool:
        """Проверка наличия изменений у компании"""
        return company_id in changes_data.get('companies', {})
    
    def _get_company_change_type(self, company_id: str, changes_data: Dict[str, Any]) -> str:
        """Получение типа изменения компании"""
        company_changes = changes_data.get('companies', {}).get(company_id, {})
        return company_changes.get('change_type', '')
    
    def _format_header_row(self, row):
        """Форматирование строки заголовков"""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        thin_border = Border(left=Side(style='thin'), 
                           right=Side(style='thin'), 
                           top=Side(style='thin'), 
                           bottom=Side(style='thin'))
        
        for cell in row:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
    
    def _format_data_row(self, row, change_type: str):
        """Форматирование строки данных в детальном отчете"""
        # Определяем цвет фона в зависимости от типа изменения
        color_map = {
            "Новые данные": "E2EFDA",  # Зеленый
            "Данные не найдены": "FCE4D6",  # Оранжевый  
            "Редактирование": "FFF2CC",  # Желтый
            "Ошибка, сайт не работает": "F8CBAD",  # Красный
            "Изменений не найдено": "D9E1F2"  # Синий
        }
        
        fill_color = color_map.get(change_type, "FFFFFF")
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        
        thin_border = Border(left=Side(style='thin'), 
                           right=Side(style='thin'), 
                           top=Side(style='thin'), 
                           bottom=Side(style='thin'))
        
        for cell in row:
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)
    
    def _format_summary_data_row(self, row, has_changes: bool, change_type: str):
        """Форматирование строки данных в сводном отчете"""
        if has_changes:
            if change_type == "Ошибка, сайт не работает":
                fill_color = "F8CBAD"  # Красный для ошибок
            else:
                fill_color = "FFF2CC"  # Желтый для изменений
        else:
            fill_color = "E2EFDA"  # Зеленый для отсутствия изменений
        
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        
        thin_border = Border(left=Side(style='thin'), 
                           right=Side(style='thin'), 
                           top=Side(style='thin'), 
                           bottom=Side(style='thin'))
        
        for cell in row:
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)
    
    def _auto_adjust_columns(self, worksheet):
        """Автоматическое регулирование ширины столбцов"""
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width